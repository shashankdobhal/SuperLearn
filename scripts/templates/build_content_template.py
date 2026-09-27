"""Regenerates docs/templates/supernova_content_template.xlsx — the content-
authoring schema described in docs/CONTENT_AUTHORING_TEMPLATE.md.

v2: every real quest slot from content/curriculum/.../days.json (all 350
days, 840 quest slots) is pre-populated in the Tasks sheet — week/day/
quest_index/quest_title/quest_type_ref/daily_mini_outcome/tasks_expected are
filled in and cell-locked, so a content writer (or an automated one) can
only ADD task rows for quest slots that actually exist, not invent a
different day-by-day structure. Week 1's 20 quest slots additionally carry
their REAL shipped content (not a placeholder) as a full worked example,
since we already have it. Sheet protection is a soft deterrent for manual
editing, not real enforcement — scripts/templates/validate_tasks_xlsx.py
is the actual backstop for anything (script- or human-authored) sent back.

Usage: python3 scripts/templates/build_content_template.py
(needs openpyxl: pip install openpyxl)
Prereq: /tmp/week1_full_dump.json — see the command in the PR/commit that
introduced this for how to regenerate it from src/content/lessons/**.
"""
import json
import os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side, Protection
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter

REPO = os.path.normpath(os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", ".."))
OUT_PATH = os.path.join(REPO, "docs", "templates", "supernova_content_template.xlsx")
WEEK1_DUMP = "/tmp/week1_full_dump.json"

wb = Workbook()

# ---------- styling helpers ----------
HEADER_FILL = PatternFill("solid", fgColor="1F2937")
HEADER_FONT = Font(name="Arial", bold=True, color="FFFFFF", size=10)
REF_FILL = PatternFill("solid", fgColor="E5E7EB")       # locked blueprint columns
EXAMPLE_FILL = PatternFill("solid", fgColor="E8F0FE")   # Week 1's real, already-shipped content
INPUT_FILL = PatternFill("solid", fgColor="FFF9DB")     # blank, ready for content
THIN = Side(style="thin", color="D0D0D0")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
BODY_FONT = Font(name="Arial", size=10)
REF_FONT = Font(name="Arial", size=10, italic=True, color="4B5563")
WRAP = Alignment(wrap_text=True, vertical="top")
LOCKED = Protection(locked=True)
UNLOCKED = Protection(locked=False)

def style_header(ws, ncols, row=1):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        cell.border = BORDER
        cell.protection = LOCKED

# ============================================================
# Load source data
# ============================================================
days_data = json.load(open(os.path.join(REPO, "content/curriculum/everyday-confidence/beginner/days.json")))
days_by_wd = {(d["week"], d["day"]): d for d in days_data}
week1 = json.load(open(WEEK1_DUMP))

REF_COLS = ["week", "day", "quest_index", "quest_title", "quest_type_ref", "daily_mini_outcome", "tasks_expected"]
CONTENT_COLS = [
    "task_type", "skill_id", "emoji", "pattern", "example", "text_hi", "text_en",
    "prompt_hi", "prompt_en", "option_1_text", "option_1_correct", "option_2_text", "option_2_correct",
    "option_3_text", "option_3_correct", "option_4_text", "option_4_correct", "correct_count_check",
    "explanation_hi", "explanation_en", "answer", "distractors", "audio_text_en",
    "is_pop_quiz", "is_final", "mission_label",
]
COLUMN_WIDTHS = {
    "week": 6, "day": 5, "quest_index": 8, "quest_title": 22, "quest_type_ref": 13,
    "daily_mini_outcome": 30, "tasks_expected": 10,
    "task_type": 10, "skill_id": 18, "emoji": 7, "pattern": 26, "example": 26,
    "text_hi": 30, "text_en": 30, "prompt_hi": 30, "prompt_en": 30,
    "option_1_text": 22, "option_1_correct": 9, "option_2_text": 22, "option_2_correct": 9,
    "option_3_text": 22, "option_3_correct": 9, "option_4_text": 22, "option_4_correct": 9,
    "correct_count_check": 9, "explanation_hi": 30, "explanation_en": 30,
    "answer": 30, "distractors": 18, "audio_text_en": 30,
    "is_pop_quiz": 9, "is_final": 8, "mission_label": 16,
}
ALL_COLS = REF_COLS + CONTENT_COLS
col_idx = {name: i + 1 for i, name in enumerate(ALL_COLS)}
N_COLS = len(ALL_COLS)

# ============================================================
# Sheet 1: Tasks
# ============================================================
ws = wb.active
ws.title = "Tasks"
ws.protection.sheet = True  # soft deterrent only — see module docstring

for i, name in enumerate(ALL_COLS, start=1):
    ws.column_dimensions[get_column_letter(i)].width = COLUMN_WIDTHS[name]
    ws.cell(row=1, column=i, value=name)
style_header(ws, N_COLS)
ws.freeze_panes = "I2"  # freeze all 7 reference columns + task_type while scrolling right
ws.row_dimensions[1].height = 30


def write_ref_cells(r, week, day, quest_index, quest_title, quest_type_ref, daily_outcome, tasks_expected):
    values = [week, day, quest_index, quest_title, quest_type_ref, daily_outcome, tasks_expected]
    for name, val in zip(REF_COLS, values):
        cell = ws.cell(row=r, column=col_idx[name], value=val)
        cell.fill = REF_FILL
        cell.font = REF_FONT
        cell.alignment = WRAP
        cell.border = BORDER
        cell.protection = LOCKED


def write_content_cell(r, name, val, fill, editable):
    cell = ws.cell(row=r, column=col_idx[name], value=val)
    cell.fill = fill
    cell.font = BODY_FONT
    cell.alignment = WRAP
    cell.border = BORDER
    cell.protection = UNLOCKED if editable else LOCKED


# Computed from col_idx rather than hardcoded — hardcoding the option-correct
# columns' letters here caused a real bug once when the reference columns
# (week/day/quest_index/...) were added in front and shifted every column
# after them; never hand-type spreadsheet cell letters that depend on layout.
_option_correct_cols = [get_column_letter(col_idx[f"option_{i}_correct"]) for i in range(1, 5)]

def correct_count_formula(r):
    return "=" + "+".join(f'COUNTIF({c}{r},"Y")' for c in _option_correct_cols)


# ---- flatten a DayLesson (from week1_full_dump.json) into content dicts ----
def flow_to_rows(day_num, skill_id, learn_flow, speak_flow):
    out = []
    for step in learn_flow:
        out.append(step_to_row(step, skill_id))
    for step in speak_flow:
        out.append(step_to_row(step, skill_id))
    return out


def step_to_row(step, skill_id):
    # skillId lives on the DayLesson, not per-step (src/lib/curriculum/
    # lesson-types.ts) — every task in a day shares one skill in this data
    # model, so it's passed in rather than read off `step`.
    row = {"task_type": step["type"] if "type" in step else "speak", "quest": step["quest"], "skill_id": skill_id}
    for k in ("emoji", "pattern", "example"):
        if k in step:
            row[k] = step[k]
    if "textHi" in step: row["text_hi"] = step["textHi"]
    if "textEn" in step: row["text_en"] = step["textEn"]
    if "promptHi" in step: row["prompt_hi"] = step["promptHi"]
    if "promptEn" in step: row["prompt_en"] = step["promptEn"]
    if "explanationHi" in step: row["explanation_hi"] = step["explanationHi"]
    if "explanationEn" in step: row["explanation_en"] = step["explanationEn"]
    if "audioTextEn" in step: row["audio_text_en"] = step["audioTextEn"]
    if "isPopQuiz" in step: row["is_pop_quiz"] = "Y" if step["isPopQuiz"] else "N"
    if "isFinal" in step: row["is_final"] = "Y" if step["isFinal"] else "N"
    if "missionLabel" in step: row["mission_label"] = step["missionLabel"]
    if "answer" in step: row["answer"] = " ".join(step["answer"])
    if "distractors" in step and step["distractors"]: row["distractors"] = ", ".join(step["distractors"])
    if "hint" in step and step["hint"]:
        row.setdefault("pattern", step["hint"]["pattern"])
        row.setdefault("example", step["hint"]["example"])
    if "options" in step:
        for i, opt in enumerate(step["options"], start=1):
            row[f"option_{i}_text"] = opt["text"]
            row[f"option_{i}_correct"] = "Y" if opt["correct"] else "N"
    return row


week1_rows_by_day = {}
for i in range(1, 8):
    d = week1[f"d{i}"]
    week1_rows_by_day[i] = flow_to_rows(i, d["skillId"], d["learnFlow"], d["speakFlow"])

r = 2
quest_slots_total = 0
week1_slots = 0
for day in days_data:
    week, day_num = day["week"], day["day"]
    quest_count = day["quest_count"]
    for qidx in range(1, quest_count + 1):
        quest_title = day.get(f"quest_{qidx}")
        quest_type = day.get(f"quest_{qidx}_type")
        quest_slots_total += 1
        write_ref_cells(r, week, day_num, qidx, quest_title, quest_type, day["daily_mini_outcome"], day["tasks_quest"])

        if week == 1:
            # Week 1 already has real, shipped content — fill it in as a
            # full worked example instead of leaving it blank.
            week1_slots += 1
            matching = [row for row in week1_rows_by_day[day_num] if row["quest"] == qidx]
            first = True
            for task_row in matching:
                if not first:
                    r += 1
                    write_ref_cells(r, week, day_num, qidx, quest_title, quest_type, day["daily_mini_outcome"], day["tasks_quest"])
                for name in CONTENT_COLS:
                    if name == "correct_count_check":
                        continue
                    write_content_cell(r, name, task_row.get(name, ""), EXAMPLE_FILL, editable=False)
                oc = col_idx["correct_count_check"]
                cell = ws.cell(row=r, column=oc,
                               value=correct_count_formula(r))
                cell.fill = EXAMPLE_FILL
                cell.font = BODY_FONT
                cell.border = BORDER
                cell.protection = LOCKED
                first = False
            r += 1
        else:
            # Blank, ready for content — one row for now; the content writer
            # duplicates it (copy row, keep columns A-G identical) for each
            # of the 5-8 tasks this quest needs.
            for name in CONTENT_COLS:
                if name == "correct_count_check":
                    continue
                write_content_cell(r, name, "", INPUT_FILL, editable=True)
            oc = col_idx["correct_count_check"]
            cell = ws.cell(row=r, column=oc,
                           value=correct_count_formula(r))
            cell.fill = INPUT_FILL
            cell.font = BODY_FONT
            cell.border = BORDER
            cell.protection = UNLOCKED
            r += 1

last_row = r - 1
print(f"quest slots: {quest_slots_total} (week 1: {week1_slots}) | data rows written: {last_row - 1}")

# ---- data validations (dropdowns) — only meaningful on unlocked rows, but
# applying to the whole column is harmless and simpler ----
dv_task_type = DataValidation(type="list", formula1='"intro,rule,mcq,build,speak"', allow_blank=True, showErrorMessage=True)
dv_task_type.error = 'Must be one of: intro, rule, mcq, build, speak'
dv_task_type.errorTitle = 'Invalid task_type'
ws.add_data_validation(dv_task_type)
dv_task_type.add(f"H2:H{last_row}")

dv_yn = DataValidation(type="list", formula1='"Y,N"', allow_blank=True, showErrorMessage=True)
dv_yn.error = 'Must be Y or N'
dv_yn.errorTitle = 'Invalid value'
ws.add_data_validation(dv_yn)
for name in ["option_1_correct", "option_2_correct", "option_3_correct", "option_4_correct", "is_pop_quiz", "is_final"]:
    col = get_column_letter(col_idx[name])
    dv_yn.add(f"{col}2:{col}{last_row}")

# ============================================================
# Sheet 2: Instructions
# ============================================================
ws2 = wb.create_sheet("Instructions")
ws2.column_dimensions["A"].width = 22
ws2.column_dimensions["B"].width = 16
ws2.column_dimensions["C"].width = 60
ws2.column_dimensions["D"].width = 45

title_font = Font(name="Arial", bold=True, size=14)
ws2["A1"] = "SuperLearn — Content Authoring Template"
ws2["A1"].font = title_font
ws2["A2"] = ("Columns A-G (grey) are pre-filled from the real curriculum blueprint for every one of the "
             "840 quest slots across all 50 weeks — don't change them, and don't add rows for a "
             "(week, day, quest_index) that isn't already here. Fill in columns H onward (yellow) only. "
             "Each quest slot starts as ONE blank row — duplicate it (copy the whole row, paste as many "
             "times as tasks_expected calls for) and fill each copy's content columns in separately; leave "
             "extra copies unused (blank task_type) if you need fewer than the max.")
ws2["A2"].font = Font(name="Arial", italic=True, size=10)
ws2["A2"].alignment = Alignment(wrap_text=True)
ws2.merge_cells("A2:D2")
ws2.row_dimensions[2].height = 60

ws2["A4"] = "Before you start"
ws2["A4"].font = Font(name="Arial", bold=True, size=12)
notes = [
    "Week 1 (rows for week=1) is filled in already with our real, shipped content — light blue, locked. "
    "Read it as the reference for tone/difficulty/progression; don't edit it.",
    "Every other quest slot (week 2 onward) starts as one blank yellow row with columns A-G already telling "
    "you exactly which week/day/quest/quest_type/daily outcome it's for — you do not need any other file to "
    "know what to write.",
    "skill_id must be one from the \"Skills Reference\" sheet. Introduce a skill only from its "
    "introduced_week onward.",
    "correct_count_check must read 1 for every mcq row (exactly one correct option) and 0 for every "
    "other row. If it doesn't, fix the option_N_correct values in that row.",
    "Don't invent a different quest structure than what's in columns A-G — if a day's real quest_count/"
    "quest types genuinely seem wrong for what you're writing, flag it as a question, don't silently work "
    "around it by adding an extra quest or skipping one.",
]
r = 5
for n in notes:
    ws2.cell(row=r, column=1, value=f"• {n}")
    ws2.cell(row=r, column=1).alignment = Alignment(wrap_text=True, vertical="top")
    ws2.merge_cells(start_row=r, start_column=1, end_row=r, end_column=4)
    ws2.row_dimensions[r].height = 40
    r += 1

r += 1
ws2.cell(row=r, column=1, value="Column reference").font = Font(name="Arial", bold=True, size=12)
r += 1
col_header_row = r
for i, h in enumerate(["Column", "Applies to", "Required?", "What it means"]):
    ws2.cell(row=col_header_row, column=i + 1, value=h)
style_header(ws2, 4, row=col_header_row)
r += 1

reference = [
    ("week / day / quest_index", "all", "pre-filled, don't edit", "Identity of the quest slot this row belongs to — matches days.json exactly."),
    ("quest_title", "all", "pre-filled, don't edit", "The quest's name from the blueprint, e.g. \"Learn the Rule\", \"Talk with Nova\"."),
    ("quest_type_ref", "all", "pre-filled, don't edit", "The quest's type from the blueprint (learn/translation/listening/writing/practice/speaking/conversation/mission/review/reading)."),
    ("daily_mini_outcome", "all", "pre-filled, don't edit", "What the learner should be able to do by the end of this day — the thing every task in this day should build toward."),
    ("tasks_expected", "all", "pre-filled, don't edit", "How many task rows this quest slot needs (always \"5-8\" in this track) — duplicate the blank row this many times."),
    ("task_type", "all", "yes", "Which of the 5 building blocks this row is: intro, rule, mcq, build, or speak. See \"Task types\" below."),
    ("skill_id", "all", "yes", "Which competency this task practices — must exist in Skills Reference, and not be introduced before its introduced_week."),
    ("emoji", "intro", "optional", "One emoji shown at the top of an intro card."),
    ("pattern", "rule, build (both required); mcq/speak (optional hint)", "conditional", "The English grammar pattern being taught, e.g. \"I live in + [Place].\" For build rows this is required (BuildStep.hint isn't optional in the app); for mcq/speak rows, filling this in (with example) shows an optional HINT box to the learner."),
    ("example", "rule, build (both required); mcq/speak (optional hint)", "conditional", "One concrete example of `pattern`, e.g. \"I live in Pune.\""),
    ("text_hi", "intro, rule", "yes for those", "The Hindi instructional copy shown on the card."),
    ("text_en", "intro", "yes", "The English instructional copy shown when the learner toggles to English (intro cards only — rule cards don't support the toggle yet)."),
    ("prompt_hi", "mcq, build, speak", "yes for those", "The Hindi-language question/prompt shown to the learner."),
    ("prompt_en", "mcq, speak", "yes for those", "The English-language version of the prompt (shown on English toggle). build has no English prompt today."),
    ("option_1..4_text", "mcq", "yes (at least 3)", "The answer choices. Use 3 or 4; leave option_4 blank if unused."),
    ("option_1..4_correct", "mcq", "yes", "Y for the single correct option, N for the rest."),
    ("correct_count_check", "mcq", "auto (formula)", "=COUNTIF of the correct columns. Must read 1 for mcq rows, 0 otherwise. Don't type into this column."),
    ("explanation_hi / explanation_en", "mcq", "yes for those", "The feedback shown after answering, in Hindi and English."),
    ("answer", "build", "yes", "The target English sentence, plain text, words separated by single spaces (e.g. \"I am from Delhi\"). Punctuation stays attached to its word (\"Delhi,\" not \"Delhi\" + \",\")."),
    ("distractors", "build", "optional", "Extra decoy words to mix into the word bank, comma-separated. Leave blank for none."),
    ("audio_text_en", "intro, mcq, build", "optional", "When filled in, the card shows a PLAY button that speaks this English text aloud, and becomes a listening-comprehension check (the learner answers/rebuilds based on what they heard, not what's printed)."),
    ("is_pop_quiz", "mcq", "optional", "Y to show the \"POP QUIZ\" banner on this question."),
    ("is_final", "speak", "yes for those", "Y on the LAST speak row of the day only — changes the header from a countdown to \"Last Question!\" (or mission_label, if set)."),
    ("mission_label", "speak", "optional", "Only on a true weekly-mission capstone task (is_final=Y): overrides \"Last Question!\" with this text, e.g. \"WEEKLY MISSION\"."),
]
for name, applies, required, desc in reference:
    ws2.cell(row=r, column=1, value=name).font = Font(name="Arial", bold=True, size=9)
    ws2.cell(row=r, column=2, value=applies)
    ws2.cell(row=r, column=3, value=required)
    ws2.cell(row=r, column=4, value=desc)
    for c in range(2, 5):
        cell = ws2.cell(row=r, column=c)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        cell.font = BODY_FONT
        cell.border = BORDER
    name_cell = ws2.cell(row=r, column=1)
    name_cell.alignment = Alignment(wrap_text=True, vertical="top")
    name_cell.border = BORDER
    ws2.row_dimensions[r].height = 34
    r += 1

r += 1
ws2.cell(row=r, column=1, value="Task types — what each one is").font = Font(name="Arial", bold=True, size=12)
r += 1
task_type_rows = [
    ("intro", "A card the learner reads (and optionally hears via audio_text_en), then taps NEXT. No question, no grading. Used to introduce a topic, or for pure listening exposure."),
    ("rule", "A card presenting one grammar pattern + one example. No question, no grading."),
    ("mcq", "A multiple-choice question: prompt + options, immediate right/wrong feedback + explanation. With audio_text_en set, it becomes a listening check (identify the exact words / the meaning / an appropriate reply)."),
    ("build", "A translate-and-build task: the learner taps English words into order to match a Hindi prompt (or, with audio_text_en set, to \"retell\" what they just heard)."),
    ("speak", "A mic-based speaking prompt. No text grading — the learner produces spoken English; the app's speaking evaluation is currently simulated (canned encouragement), not real AI grading yet."),
]
for name, desc in task_type_rows:
    ws2.cell(row=r, column=1, value=name).font = Font(name="Arial", bold=True, size=9)
    ws2.cell(row=r, column=2, value=desc)
    ws2.merge_cells(start_row=r, start_column=2, end_row=r, end_column=4)
    ws2.cell(row=r, column=2).alignment = Alignment(wrap_text=True, vertical="top")
    ws2.row_dimensions[r].height = 34
    r += 1

r += 1
ws2.cell(row=r, column=1, value="Pedagogical guardrails (please follow, don't just fill cells)").font = Font(name="Arial", bold=True, size=12)
r += 1
guardrails = [
    "Every week needs ONE weekly speaking outcome; every day needs one small mini-outcome that visibly "
    "builds toward it (see docs/CURRICULUM_PHILOSOPHY.md sections 7-8 if you have that file).",
    "New vocabulary/grammar should connect to what a PRIOR day already taught, not appear cold — reuse "
    "names, places and topics across a week where it makes sense, the way Days 1-7 of Week 1 do.",
    "Don't reuse the exact same prompt text/instruction across many different days (e.g. a listening "
    "prompt that reads identically on every single day is a sign something's wrong, not a template to keep).",
    "Don't repeat the exact same example sentence many times across the track — vary the surface content "
    "even when the grammar pattern repeats.",
    "Every day should have at least one speak task; avoid days with zero spoken output.",
    "If in doubt about whether a day's plan actually produces more real speaking (per "
    "CURRICULUM_PHILOSOPHY.md's core invariant), that question matters more than filling every column.",
]
for g in guardrails:
    ws2.cell(row=r, column=1, value=f"• {g}")
    ws2.merge_cells(start_row=r, start_column=1, end_row=r, end_column=4)
    ws2.cell(row=r, column=1).alignment = Alignment(wrap_text=True, vertical="top")
    ws2.row_dimensions[r].height = 32
    r += 1

# ============================================================
# Sheet 3: Skills Reference (mirror of skills.json)
# ============================================================
ws3 = wb.create_sheet("Skills Reference")
skills = json.load(open(os.path.join(REPO, "content/curriculum/everyday-confidence/beginner/skills.json")))
skill_headers = ["skill_id", "skill", "role", "introduced_week", "suggested_prerequisite_relationship"]
for i, h in enumerate(skill_headers, start=1):
    ws3.cell(row=1, column=i, value=h)
    ws3.column_dimensions[get_column_letter(i)].width = 26 if i != 4 else 16
style_header(ws3, len(skill_headers))
for r2, s in enumerate(skills, start=2):
    ws3.cell(row=r2, column=1, value=s["skill_id"])
    ws3.cell(row=r2, column=2, value=s["skill"])
    ws3.cell(row=r2, column=3, value=s["role"])
    ws3.cell(row=r2, column=4, value=s["introduced_week"])
    ws3.cell(row=r2, column=5, value=s["suggested_prerequisite_relationship"])
    for c in range(1, 6):
        cell = ws3.cell(row=r2, column=c)
        cell.font = BODY_FONT
        cell.alignment = WRAP
        cell.border = BORDER
ws3.freeze_panes = "A2"
note_row = len(skills) + 3
ws3.cell(row=note_row, column=1, value=(
    "Note: 'describe_place' appears twice in the source data (week 2 and week 39, slightly different "
    "prerequisite notes) — a source data quirk, not fixed here. Pick whichever occurrence's week fits."
))
ws3.cell(row=note_row, column=1).font = Font(name="Arial", italic=True, size=9)
ws3.merge_cells(start_row=note_row, start_column=1, end_row=note_row, end_column=5)
ws3.cell(row=note_row, column=1).alignment = Alignment(wrap_text=True)

# ============================================================
# Sheet 4: Blueprint Reference (mirror of days.json, all 350 days)
# ============================================================
ws4 = wb.create_sheet("Blueprint Reference")
bp_headers = ["week", "day", "week_arc", "weekly_outcome", "daily_mini_outcome", "activity_mix",
              "quest_count", "quest_1", "quest_1_type", "quest_2", "quest_2_type",
              "quest_3", "quest_3_type", "quest_4", "quest_4_type", "primary_skills"]
for i, h in enumerate(bp_headers, start=1):
    ws4.cell(row=1, column=i, value=h)
    ws4.column_dimensions[get_column_letter(i)].width = 16 if h not in ("weekly_outcome", "daily_mini_outcome") else 30
style_header(ws4, len(bp_headers))
for r2, d in enumerate(days_data, start=2):
    for i, h in enumerate(bp_headers, start=1):
        ws4.cell(row=r2, column=i, value=d.get(h))
    for c in range(1, len(bp_headers) + 1):
        cell = ws4.cell(row=r2, column=c)
        cell.font = BODY_FONT
        cell.alignment = WRAP
        cell.border = BORDER
ws4.freeze_panes = "A2"

os.makedirs(os.path.dirname(OUT_PATH), exist_ok=True)
wb.save(OUT_PATH)
print(f"saved: {OUT_PATH}")
